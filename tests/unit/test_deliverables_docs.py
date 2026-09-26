"""Deliverables: documents rendered from vault notes, drafted through the
production pipeline, gated by Aaron's client_sensitive approval."""

from __future__ import annotations

import asyncio
import importlib
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

ADMIN = {"x-api-key": "admin-key"}
READ = {"x-api-key": "test-key"}

SAMPLE = """# Retention in first year

## Summary

- Online first-years leave **earlier** than on-campus peers [1].
- The gap opens in weeks 2 to 4.

## Findings

1. Part-time study
2. Later start in life
    - Many are carers

| Measure | Online | On campus |
| --- | --- | --- |
| Part-time share | 68% | 22% |

> Retention is an operating-model problem.

## Sources

1. Department of Education, 2024.
"""


# ---------------------------------------------------------------------------
# Renderer
# ---------------------------------------------------------------------------

def test_docx_render_brand_structure_and_no_overwrite(tmp_path):
    from docx import Document

    from media import doc_render

    out = tmp_path / "brief-v1.docx"
    result = doc_render.render_docx(SAMPLE, out, {"title": "Retention in first year",
                                                  "type_label": "Client briefing", "version": 1})
    assert result["bytes"] > 0 and out.is_file()
    doc = Document(str(out))
    names = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Title", "Retention in first year") in names
    # The note's H1 is the title on the cover, so sections become Heading 1.
    assert ("Heading 1", "Summary") in names and ("Heading 1", "Findings") in names
    assert sum(1 for n, _ in names if n == "Heading 1" and _ == "Retention in first year") == 0
    assert any(n == "List Bullet" and "earlier" in t for n, t in names)
    assert any(t.startswith("1.") and "Part-time study" in t for _, t in names)
    assert any(n == "Quote" for n, _ in names)
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Measure", "Online", "On campus"]
    assert doc.styles["Normal"].font.name == "Open Sans"
    assert doc.styles["Heading 1"].font.name == "Open Sans"
    header = " ".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "DRAFT FOR REVIEW" in header
    with pytest.raises(FileExistsError):
        doc_render.render_docx(SAMPLE, out, {"title": "again"})


def test_working_notes_never_render(tmp_path):
    from media import doc_render
    from orchestrator.uploads import extract_text

    body = doc_render.join_working_notes(SAMPLE, "### Quality review\n\nSECRET-REVIEWER-TEXT")
    doc_md, notes = doc_render.split_working_notes(body)
    assert "SECRET-REVIEWER-TEXT" in notes and "SECRET" not in doc_md
    out = tmp_path / "x-v1.docx"
    doc_render.render_docx(body, out, {"title": "Retention in first year", "status": "final"})
    text = extract_text("x-v1.docx", out.read_bytes())
    assert "SECRET-REVIEWER-TEXT" not in text
    assert "Summary" in text and "Findings" in text and "Sources" in text


def test_pptx_render_has_notes_on_every_slide(tmp_path):
    pptx = pytest.importorskip("pptx")
    from media import doc_render

    deck = "# Deck\nSubtitle here\n\n" + "\n\n".join(
        f"## Slide {i}: Point {i}\n- point a\n- point b\nNotes: say thing {i}." for i in range(1, 12)
    )
    assert doc_render.parse_slides(deck)["slides"][0]["title"] == "Point 1"
    out = tmp_path / "deck-v1.pptx"
    result = doc_render.render_pptx(deck, out, {"title": "Deck"})
    assert result["slides"] == 12
    prs = pptx.Presentation(str(out))
    assert len(prs.slides) == 12
    for slide in prs.slides:
        assert slide.notes_slide.notes_text_frame.text.strip()


def test_agent_markdown_is_cleaned_to_house_style():
    from orchestrator import deliverables

    raw = ("Sure, here it is.\n\n## Background\nOnline share rose \u2014 sharply, then fell \u2013 briefly, in 2024\u20132025.\n"
           "[To confirm: 2025 figure]\n[To confirm: owner]\n\n## Sources\n[1] DoE 2024\n[2] QILT 2024\n"
           "| a | b |\n| --- | --- |\n| [1] | x |\n")
    md = deliverables._clean_markdown(raw, "Paper")
    assert md.startswith("# Paper\n")          # title added, preamble dropped
    assert "Sure, here it is" not in md
    assert "\u2014" not in md and "rose, sharply" in md
    assert "fell, briefly" in md and "2024\u20132025" in md  # spaced en dash goes, ranges stay
    assert "- [To confirm: 2025 figure]" in md and "- [1] DoE 2024" in md
    assert "| [1] | x |" in md                  # tables untouched


def test_agent_json_survives_preamble_and_fences():
    from orchestrator import deliverables

    reply = 'Here is my review.\n\n```json\n{"verdict": "revise", "checks": {"costs": "pass"}}\n```\nThanks.'
    assert deliverables._parse_json(reply)["verdict"] == "revise"
    assert deliverables._parse_json('{"verdict": "pass"}')["verdict"] == "pass"
    assert "_raw" in deliverables._parse_json("no json at all")


def test_section_budgets_add_up_to_the_type_target():
    from orchestrator import deliverables
    from orchestrator.doc_types import DOC_TYPES

    for key in ("client_briefing", "decision_paper", "proposal"):
        spec = DOC_TYPES[key]
        plan = deliverables._section_plan({"script": {"outline": {"sections": [
            {"heading": "Options", "points": ["compare three models"]}]}}}, spec)
        low, high = deliverables._word_range(spec)
        total = sum(s["words"] for s in plan)
        assert "Sources" not in [s["heading"] for s in plan]
        assert low <= total <= high, (key, total)
        if key == "decision_paper":
            options = next(s for s in plan if s["heading"] == "Options")
            assert options["points"] == ["compare three models"]
            assert options["words"] > next(s for s in plan if s["heading"] == "Decision sought")["words"]


def test_office_convert_is_off_when_disabled(tmp_path, monkeypatch):
    from media import office_convert

    monkeypatch.setenv("DELIVERABLES_PDF", "off")
    src = tmp_path / "a.docx"
    src.write_bytes(b"x")
    result = office_convert.convert(src)
    assert result["ok"] is False and result.get("skipped")


# ---------------------------------------------------------------------------
# API and pipeline
# ---------------------------------------------------------------------------

def _client(tmp_path, monkeypatch):
    monkeypatch.setenv("MEDIA_DB_PATH", str(tmp_path / "media.db"))
    monkeypatch.setenv("MEDIA_DERIVED_ROOT", str(tmp_path / "derived"))
    monkeypatch.setenv("REMOTION_DIR", str(tmp_path / "missing-remotion"))
    monkeypatch.setenv("OBSIDIAN_VAULT_PATH", str(tmp_path / "vault"))
    monkeypatch.setenv("API_KEY", "test-key")
    monkeypatch.setenv("ADMIN_API_KEY", "admin-key")
    monkeypatch.setenv("DELIVERABLES_RETRIEVAL", "off")
    monkeypatch.setenv("DELIVERABLES_PDF", "off")
    monkeypatch.setenv("DELIVERABLES_SYNC", "1")
    monkeypatch.setenv("DELIVERABLES_NOTIFY", "off")
    from orchestrator import dashboard, deliverables, governance, operating, production
    import orchestrator.main as main

    importlib.reload(dashboard)
    importlib.reload(governance)
    importlib.reload(operating)
    importlib.reload(production)
    importlib.reload(deliverables)
    main = importlib.reload(main)
    return TestClient(main.app), main, production, deliverables


DECK_MD = "# {title}\nFor the executive team\n\n" + "\n\n".join(
    f"## Point {i}\n- evidence {i}\nNotes: explain point {i}." for i in range(1, 9)
)


def _filler(words: int) -> str:
    sentence = "Online cohorts differ from campus cohorts in timing [1]."
    return " ".join([sentence] * max(1, words // 8 + 1))


def _fake_agent(calls):
    async def fake(department, subagent, query, rag_context=None, **kwargs):
        import re as _re
        calls.append((department, subagent, query))
        section = _re.search(r'Write only the "## ([^"]+)" section', query)
        if "needs about" in query and "This draft of the" in query:
            heading = _re.search(r'This draft of the "## ([^"]+)" section', query).group(1)
            target = int(_re.search(r"needs about (\d+)", query).group(1))
            return {"answer": f"## {heading}\n\n{_filler(target)}", "cost_usd": 0.001}
        if section:
            heading = section.group(1)
            target = int(_re.search(r"About (\d+) words", query).group(1))
            body = "Too short here [1]." if heading == "Background" else _filler(target)
            if heading == "Options":
                body += "\n\n| Option | Benefit |\n| --- | --- |\n| Hybrid | Control [1] |"
            if heading == "Summary" or heading == "Decision sought":
                body += " It is an operating \u2014 model question."
            if heading == "Next steps" and "proposal" in query:
                body += " Our day rate is $2,000."
            return {"answer": f"Sure.\n\n## {heading}\n\n{body}\n\n## Sources\n\n- [1] model-added", "cost_usd": 0.001}
        if "structured brief for the writer" in query:
            answer = json.dumps({"purpose": "Inform", "key_questions": ["Why?"], "sections": []})
        elif "list the evidence" in query:
            answer = json.dumps({"facts": [{"claim": "68% part-time", "source": "[1]"}],
                                 "sources": [{"ref": "[1]", "title": "DoE 2024"},
                                             {"ref": "[2]", "title": "--- cap: junk frontmatter ---"}],
                                 "gaps": ["No 2025 data"]})
        elif "Outline the document" in query or "Plan a deck" in query:
            answer = json.dumps({"sections": [{"heading": "Summary", "points": ["gap"]}]})
        elif "Write the deck" in query:
            title = query.split("Title: ", 1)[1].split("\n", 1)[0]
            answer = DECK_MD.format(title=title)
        elif "Return only the markdown document" in query:
            title = query.split("Title: ", 1)[1].split("\n", 1)[0]
            body = SAMPLE.replace("# Retention in first year", f"# {title}")
            if "Proposal" in query:
                body += "\n## Next steps\n\nOur day rate is $2,000.\n"
            answer = "Here is the document:\n\n" + body.replace("operating-model", "operating — model")
        elif query.startswith("Edit this") or query.startswith("Tighten this"):
            answer = query.split("\n\n", 2)[-1]
        elif query.startswith("Review this"):
            answer = json.dumps({"verdict": "pass", "summary": "Solid.",
                                 "checks": {"evidence": "pass: cited", "no_fees": "pass"},
                                 "issues": ["Add 2025 data"], "to_confirm": []})
        else:
            answer = "Handoff: use in the Provost meeting."
        return {"answer": answer, "cost_usd": 0.001}
    return fake


def _create(client, deliverables, monkeypatch, calls, doc_type="decision_paper", **extra):
    monkeypatch.setattr(deliverables, "_call_agent", _fake_agent(calls))
    payload = {"doc_type": doc_type, "title": "Online model options for council",
               "brief": "Recommend an online operating model for council in November.",
               "client": "Swinburne Online", "audience": "Academic council", "run": False, **extra}
    created = client.post("/deliverables", headers=ADMIN, json=payload)
    assert created.status_code == 200, created.text
    did = created.json()["item"]["id"]
    result = asyncio.run(deliverables.run_pipeline(did))
    assert result["ok"], result
    return did


def test_list_is_real_notes_only(tmp_path, monkeypatch):
    client, *_ = _client(tmp_path, monkeypatch)
    empty = client.get("/deliverables", headers=READ).json()
    assert empty["items"] == [] and empty["demo"] is False
    demo = client.get("/deliverables?demo=1", headers=READ).json()
    assert demo["demo"] is True and demo["items"]
    monkeypatch.delenv("OBSIDIAN_VAULT_PATH")
    monkeypatch.setattr("orchestrator.dashboard._vault_root", lambda: None)
    missing = client.get("/deliverables", headers=READ).json()
    assert missing["items"] == [] and "not configured" in missing["error"]


def test_document_runs_to_review_with_docx_v1(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    calls: list = []
    did = _create(client, deliverables, monkeypatch, calls)

    detail = client.get(f"/deliverables/{did}", headers=READ).json()
    assert detail["production"]["state"] == "review"
    # Sectioned drafting: every required section, in order, to its budget.
    from media.doc_render import word_count
    body = detail["body"]
    assert word_count(body) >= 1400  # decision paper target is 1,400 to 2,800
    order = [body.index(f"## {h}") for h in ("Decision sought", "Background", "Options", "Recommendation",
                                              "Implementation", "Risks", "Sources")]
    assert order == sorted(order)
    assert body.count("## Sources") == 1 and "- [1] DoE 2024" in body and "model-added" not in body
    assert "junk frontmatter" not in body  # uncited source [2] is left out
    later = [q for d, s, q in calls if 'Write only the "## Background"' in q]
    assert later and "Already written above" in later[0] and "## Decision sought" in later[0]
    assert "Sure." not in body and "Too short here" not in body  # preamble dropped, short section expanded
    assert detail["item"]["words"] >= 1400
    assert any("1 expanded" in (e["note"] or "") for e in detail["production"]["events"])
    assert detail["item"]["readiness"] == "Awaiting approval"
    assert detail["item"]["readiness"] != "Client-ready"
    assert detail["versions"][0]["version"] == 1 and "docx" in detail["versions"][0]
    assert detail["body"].startswith("# Online model options for council")
    assert "—" not in detail["body"]  # house style applied
    assert "Quality review by operations/quality-reviewer" in detail["working_notes"]
    assert "Evidence gaps" in detail["working_notes"]
    assert detail["item"]["cost_usd"] > 0
    notes = [e["note"] for e in detail["production"]["events"]]
    assert any("asset_plan skipped" in n for n in notes)
    states = [e["to_state"] for e in detail["production"]["events"]]
    assert states == ["brief", "research", "outline", "draft", "render", "review"]
    used = {(d, s) for d, s, _ in calls}
    assert ("research_intelligence", "insights-strategist") in used
    assert ("operations", "quality-reviewer") in used

    file = client.get(f"/deliverables/{did}/file/1?fmt=docx", headers=READ)
    assert file.status_code == 200 and file.content[:2] == b"PK"

    listed = client.get("/deliverables", headers=READ).json()["items"]
    assert listed[0]["id"] == did and listed[0]["readiness"] == "Awaiting approval"


def test_rejection_then_approval_finalises(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    did = _create(client, deliverables, monkeypatch, [])
    pid = client.get(f"/deliverables/{did}", headers=READ).json()["production"]["production_id"]

    pending = client.get("/governance/pending", headers=READ).json()["items"]
    assert [p["gate"] for p in pending if p["production_id"] == pid] == ["client_sensitive"]

    rejected = client.post("/governance/approve", headers=ADMIN, json={
        "gate": "client_sensitive", "target_id": pid, "actor": "Aaron", "note": "Tighten options", "status": "rejected"})
    assert rejected.status_code == 200
    detail = client.get(f"/deliverables/{did}", headers=READ).json()
    assert detail["item"]["readiness"] == "Changes requested"
    assert "Tighten options" in detail["working_notes"]

    approved = client.post("/governance/approve", headers=ADMIN, json={
        "gate": "client_sensitive", "target_id": pid, "actor": "Aaron", "note": "", "status": "approved"})
    assert approved.status_code == 200
    detail = client.get(f"/deliverables/{did}", headers=READ).json()
    assert detail["production"]["state"] == "publish"
    assert detail["item"]["readiness"] == "Client-ready"
    assert detail["item"]["approved_by"] == "Aaron"
    assert [v["version"] for v in detail["versions"]] == [2, 1]

    from docx import Document
    vault = tmp_path / "vault" / "13_Command Centre" / "Deliverables" / "_files" / did
    header = " ".join(p.text for p in Document(str(vault / f"{did}-v2.docx")).sections[0].header.paragraphs)
    assert "CONFIDENTIAL" in header

    locked = client.post(f"/deliverables/{did}/render", headers=ADMIN)
    assert locked.status_code == 409


def test_rerender_after_note_edit_keeps_v1(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    did = _create(client, deliverables, monkeypatch, [])
    note = tmp_path / "vault" / "13_Command Centre" / "Deliverables" / f"{did}.md"
    note.write_text(note.read_text(encoding="utf-8").replace("Online cohorts differ", "EDITED-IN-OBSIDIAN Online cohorts differ", 1),
                    encoding="utf-8")
    rendered = client.post(f"/deliverables/{did}/render", headers=ADMIN)
    assert rendered.status_code == 200, rendered.text
    assert rendered.json()["version"] == 2
    folder = note.parent / "_files" / did
    assert (folder / f"{did}-v1.docx").is_file() and (folder / f"{did}-v2.docx").is_file()
    from orchestrator.uploads import extract_text
    assert "EDITED-IN-OBSIDIAN" in extract_text("v2.docx", (folder / f"{did}-v2.docx").read_bytes())
    assert "EDITED-IN-OBSIDIAN" not in extract_text("v1.docx", (folder / f"{did}-v1.docx").read_bytes())


def test_drawer_action_is_saved_in_working_notes(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    calls: list = []
    did = _create(client, deliverables, monkeypatch, calls)
    out = client.post(f"/deliverables/{did}/action", headers=ADMIN, json={"action": "summary"})
    assert out.status_code == 200 and out.json()["saved"] is True
    assert "Document:\n# Online model options" in calls[-1][2]
    notes = client.get(f"/deliverables/{did}", headers=READ).json()["working_notes"]
    assert "Handoff summary by support/responder-agent" in notes
    assert "Provost meeting" in notes
    bad = client.post(f"/deliverables/{did}/action", headers=ADMIN, json={"action": "send"})
    assert bad.status_code == 422


def test_proposal_fee_mentions_are_flagged(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    calls: list = []
    did = _create(client, deliverables, monkeypatch, calls, doc_type="proposal")
    notes = client.get(f"/deliverables/{did}", headers=READ).json()["working_notes"]
    assert "Check before approval" in notes and "day rate" in notes
    assert ("marketing_sales", "copywriter") in {(d, s) for d, s, _ in calls}


def test_deck_is_built_from_a_document(tmp_path, monkeypatch):
    pptx = pytest.importorskip("pptx")
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    source = _create(client, deliverables, monkeypatch, [])
    calls: list = []
    monkeypatch.setattr(deliverables, "_call_agent", _fake_agent(calls))
    no_source = client.post("/deliverables", headers=ADMIN, json={"doc_type": "slide_deck", "title": "Council deck", "run": False})
    assert no_source.status_code == 422
    created = client.post("/deliverables", headers=ADMIN, json={
        "doc_type": "slide_deck", "title": "Council deck", "source_deliverable_id": source, "run": False})
    assert created.status_code == 200, created.text
    did = created.json()["item"]["id"]
    assert asyncio.run(deliverables.run_pipeline(did))["ok"]
    detail = client.get(f"/deliverables/{did}", headers=READ).json()
    assert detail["production"]["state"] == "review"
    assert "pptx" in detail["versions"][0]
    path = tmp_path / "vault" / "13_Command Centre" / "Deliverables" / "_files" / did / f"{did}-v1.pptx"
    prs = pptx.Presentation(str(path))
    assert len(prs.slides) == 9
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
    assert any("Online model options for council" in q for _, _, q in calls)


def test_governance_documents_need_only_client_sensitive(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    from orchestrator import governance
    doc = production.create_production("Paper", "demo", "decision_paper")
    vid = production.create_production("Clip", "demo", "linkedin_short")
    production.transition(doc, "review", "t")
    production.transition(vid, "review", "t")
    assert governance.required_for_transition(production.get_production(doc), "publish") == ["client_sensitive"]
    assert governance.required_for_transition(production.get_production(doc), "review") == []
    assert set(governance.required_for_transition(production.get_production(vid), "publish")) >= {
        "public_claim", "external_publish"}


def test_daemon_does_not_mint_tasks_for_documents(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    from orchestrator import operating
    production.create_production("Paper", "demo", "decision_paper")
    vid = production.create_production("Clip", "demo", "linkedin_short")
    created = operating.sync_production_tasks()
    assert [c["production_id"] for c in created] == [vid]


def test_video_advance_still_uses_remotion_path(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    from media import render as render_service
    pid = production.create_production("Clip", "demo", "linkedin_short")
    production.transition(pid, "asset_plan", "t")
    seen = []

    async def fake_agent(**kwargs):
        return {"answer": "{}"}

    def fake_render(production_id, template, props):
        seen.append(template)
        return {"asset_id": None}

    monkeypatch.setattr(production, "call_wijerco_agent", fake_agent)
    monkeypatch.setattr(render_service, "render", fake_render)
    asyncio.run(production.advance(pid))
    assert seen == ["linkedin_short"]
    assert production.get_production(pid)["state"] == "render"


def test_archive_moves_note_and_cancels_production(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    did = _create(client, deliverables, monkeypatch, [])
    pid = client.get(f"/deliverables/{did}", headers=READ).json()["production"]["production_id"]
    out = client.post(f"/deliverables/{did}/archive", headers=ADMIN)
    assert out.status_code == 200
    assert out.json()["archived_to"].endswith(f"_archive/{did}.md")
    assert client.get(f"/deliverables/{did}", headers=READ).status_code == 404
    assert production.get_production(pid)["state"] == "cancelled"
    assert client.get("/deliverables", headers=READ).json()["items"] == []


def test_handoff_updates_the_existing_document_note(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    did = _create(client, deliverables, monkeypatch, [])
    pid = client.get(f"/deliverables/{did}", headers=READ).json()["production"]["production_id"]
    client.post("/governance/approve", headers=ADMIN, json={
        "gate": "client_sensitive", "target_id": pid, "actor": "Aaron", "note": "", "status": "approved"})
    out = client.post(f"/production/{pid}/handoff", headers=ADMIN, json={"actor": "Aaron", "note": "Sent Monday"})
    assert out.status_code == 200, out.text
    assert out.json()["deliverable"]["updated"] is True
    notes = list((tmp_path / "vault" / "13_Command Centre" / "Deliverables").glob("*.md"))
    assert len(notes) == 1


def test_bad_requests(tmp_path, monkeypatch):
    client, main, production, deliverables = _client(tmp_path, monkeypatch)
    assert client.post("/deliverables", headers=ADMIN, json={"doc_type": "memo", "title": "x" * 5}).status_code == 422
    assert client.post("/deliverables", headers=ADMIN, json={
        "doc_type": "proposal", "title": "Short", "brief": "hi"}).status_code == 422
    assert client.post("/deliverables", headers={"x-api-key": "wrong"}, json={
        "doc_type": "proposal", "title": "Short", "brief": "a real brief here", "run": False}).status_code in (401, 403)
    assert client.get("/deliverables/..%5Csecret", headers=READ).status_code in (400, 404)
    types = client.get("/deliverables/types", headers=READ).json()["items"]
    assert {t["key"] for t in types} == {"client_briefing", "decision_paper", "proposal", "slide_deck"}
