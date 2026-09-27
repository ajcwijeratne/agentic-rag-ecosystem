"""Render Deliverables notes to branded Office files.

The note body in the vault is the source of truth for a document. This module
turns that markdown into a WijerCo-styled DOCX (python-docx) or PPTX
(python-pptx). It never edits the note and never overwrites an existing file:
callers pass a fresh versioned path each time.

Both renderers import their Office library lazily. wijerco's deploy script
import-checks `orchestrator.main` before it restarts anything, so a missing
package must fail a render, not the import.
"""

from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path
from typing import Any

FONT = "Open Sans"
MONO = "Consolas"
PINE = (0x1F, 0x4D, 0x3F)
GOLD = (0xC2, 0x8F, 0x1E)
INK = (0x1D, 0x26, 0x23)
MUTED = (0x5B, 0x6B, 0x66)
RULE = "C9D2CE"

# Anything after this marker in a note is working notes (reviews, agent
# output, reminders). It is shown in Obsidian and the drawer, never rendered.
WORKING_NOTES_MARKER = "%% wijerco:working-notes %%"
PAGEBREAK_MARKERS = {"<!-- pagebreak -->", "\\pagebreak", "[pagebreak]"}

REPO_ROOT = Path(__file__).resolve().parents[1]
DOCX_TEMPLATE = Path(os.getenv("DOC_TEMPLATE_DOCX", str(REPO_ROOT / "templates" / "docs" / "wijerco-reference.docx")))

_SAFE_LINK = re.compile(r"^(https?://|mailto:)", re.I)


# ---------------------------------------------------------------------------
# Note body helpers
# ---------------------------------------------------------------------------

def split_working_notes(body: str) -> tuple[str, str]:
    """Return (document markdown, working notes markdown)."""
    text = body or ""
    idx = text.find(WORKING_NOTES_MARKER)
    if idx == -1:
        return text.rstrip() + ("\n" if text.strip() else ""), ""
    return text[:idx].rstrip() + "\n", text[idx + len(WORKING_NOTES_MARKER):].strip("\n")


def join_working_notes(document_md: str, notes_md: str) -> str:
    doc = (document_md or "").rstrip() + "\n"
    notes = (notes_md or "").strip("\n")
    if not notes:
        return doc
    return f"{doc}\n{WORKING_NOTES_MARKER}\n\n{notes}\n"


def word_count(markdown: str) -> int:
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'’-]*", markdown or ""))


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def strip_inline_md(s: str) -> str:
    s = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", s or "")
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"(\*\*|__)(.+?)\1", r"\2", s)
    s = re.sub(r"(?<![\w*])[*_](.+?)[*_](?![\w*])", r"\1", s)
    return s.replace("`", "").strip()


# ---------------------------------------------------------------------------
# DOCX: styles and template
# ---------------------------------------------------------------------------

def _rgb(t):
    from docx.shared import RGBColor
    return RGBColor(*t)


def _force_font(rpr_owner, font: str = FONT) -> None:
    """Set explicit fonts on an rPr and drop theme font references, which
    otherwise win over the explicit name in Word."""
    from docx.oxml.ns import qn

    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)


def _remove_child(el, tag: str) -> None:
    from docx.oxml.ns import qn

    if el is None:
        return
    found = el.find(qn(tag))
    if found is not None:
        el.remove(found)


def apply_brand(doc) -> None:
    """Configure page size, fonts, colours and the styles the renderer uses."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    # Document defaults: every run falls back to Open Sans, not the theme font.
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is not None:
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None:
                    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                        rfonts.attrib.pop(qn(attr), None)
                    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                        rfonts.set(qn(attr), FONT)

    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE):
            el = style.element
            rpr = el.find(qn("w:rPr"))
            if rpr is not None and rpr.find(qn("w:rFonts")) is not None:
                _force_font(el)

    def para(name: str):
        try:
            return doc.styles[name]
        except KeyError:
            return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    normal = para("Normal")
    _force_font(normal.element)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    specs = {
        "Title": dict(size=30, bold=True, color=PINE, before=0, after=8),
        "Subtitle": dict(size=14, bold=False, color=MUTED, before=0, after=18, italic=False),
        "Heading 1": dict(size=16, bold=True, color=PINE, before=18, after=6),
        "Heading 2": dict(size=13, bold=True, color=PINE, before=14, after=4),
        "Heading 3": dict(size=11, bold=True, color=INK, before=10, after=3),
    }
    for name, s in specs.items():
        st = para(name)
        _force_font(st.element)
        _remove_child(st.element.pPr, "w:pBdr")
        st.font.size = Pt(s["size"])
        st.font.bold = s["bold"]
        st.font.italic = s.get("italic", False)
        st.font.color.rgb = _rgb(s["color"])
        st.paragraph_format.space_before = Pt(s["before"])
        st.paragraph_format.space_after = Pt(s["after"])
        if name.startswith("Heading"):
            st.paragraph_format.keep_with_next = True

    quote = para("Quote")
    _force_font(quote.element)
    quote.font.italic = True
    quote.font.color.rgb = _rgb(MUTED)
    quote.paragraph_format.left_indent = Cm(0.8)

    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Paragraph"):
        try:
            st = doc.styles[name]
            _force_font(st.element)
            st.paragraph_format.space_after = Pt(3)
        except KeyError:
            pass

    # The built-in Header and Footer styles carry Letter-width tab stops that
    # would catch our right-aligned tab first. Drop them.
    for name in ("Header", "Footer"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _force_font(st.element)
        _remove_child(st.element.pPr, "w:tabs")

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)


def build_reference_template(path: Path = DOCX_TEMPLATE) -> Path:
    """Write the WijerCo reference template. Open it in Word to adjust styles;
    the renderer picks up whatever the template defines."""
    from docx import Document

    doc = Document()
    apply_brand(doc)
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _new_document():
    from docx import Document

    if DOCX_TEMPLATE.is_file():
        doc = Document(str(DOCX_TEMPLATE))
        body = doc.element.body
        for child in list(body):
            if not child.tag.endswith("}sectPr"):
                body.remove(child)
    else:
        doc = Document()
    apply_brand(doc)
    return doc


# ---------------------------------------------------------------------------
# DOCX: low-level helpers
# ---------------------------------------------------------------------------

def _style(doc, name: str, fallback: str = "Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles[fallback]


def _border(paragraph, color: str = RULE, size: int = 6, edge: str = "bottom") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), str(size))
    line.set(qn("w:space"), "4")
    line.set(qn("w:color"), color)
    bdr.append(line)
    ppr.insert_element_before(
        bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
        "w:pPrChange",
    )


def _shade_cell(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                               "w:vAlign", "w:hideMark")


def _table_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), RULE)
        borders.append(el)
    tblpr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement

    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))


def _field(paragraph, instr: str, size=None, color=None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run_with(el):
        run = paragraph.add_run()
        if size:
            run.font.size = size
        if color:
            run.font.color.rgb = _rgb(color)
        run._r.append(el)
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run_with(instr_el)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = paragraph.add_run("1")
    if size:
        r.font.size = size
    if color:
        r.font.color.rgb = _rgb(color)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def _hyperlink(paragraph, url: str, text: str, bold: bool, italic: bool, size=None) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "%02X%02X%02X" % PINE)
    rpr.append(color)
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        rpr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _inline(paragraph, inline_token, *, bold=False, italic=False, color=None, size=None) -> None:
    """Write a markdown-it inline token's children as runs."""
    from docx.shared import Pt

    state = {"bold": bold, "italic": italic, "strike": False}
    link: dict[str, Any] | None = None

    def emit(text: str, mono: bool = False) -> None:
        if not text:
            return
        if link is not None:
            link["parts"].append(text)
            return
        run = paragraph.add_run(text)
        run.bold = state["bold"] or None
        run.italic = state["italic"] or None
        if state["strike"]:
            run.font.strike = True
        if mono:
            run.font.name = MONO
            run.font.size = Pt(9.5)
        elif size is not None:
            run.font.size = size
        if color is not None:
            run.font.color.rgb = _rgb(color)

    for child in (inline_token.children or []):
        kind = child.type
        if kind == "text":
            emit(child.content)
        elif kind == "softbreak":
            emit(" ")
        elif kind == "hardbreak":
            if link is None:
                paragraph.add_run().add_break()
        elif kind == "strong_open":
            state["bold"] = True
        elif kind == "strong_close":
            state["bold"] = bold
        elif kind == "em_open":
            state["italic"] = True
        elif kind == "em_close":
            state["italic"] = italic
        elif kind == "s_open":
            state["strike"] = True
        elif kind == "s_close":
            state["strike"] = False
        elif kind == "code_inline":
            emit(child.content, mono=True)
        elif kind == "link_open":
            link = {"href": str(child.attrs.get("href") or ""), "parts": []}
        elif kind == "link_close":
            current, link = link, None
            if current is None:
                continue
            text = "".join(current["parts"]) or current["href"]
            if _SAFE_LINK.match(current["href"]):
                _hyperlink(paragraph, current["href"], text, state["bold"], state["italic"], size)
            else:
                emit(text)
        elif kind == "image":
            emit(f"[Image: {child.content or 'figure'}]")
        elif kind in ("html_inline",):
            emit(child.content)


# ---------------------------------------------------------------------------
# DOCX: markdown walk
# ---------------------------------------------------------------------------

def _markdown_it():
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark", {"html": False}).enable("table").enable("strikethrough")


def _add_markdown(doc, markdown: str, skip_title: str | None = None) -> None:
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt

    tokens = _markdown_it().parse(markdown or "")
    # When the note's only H1 is the title (shown on the cover), promote the
    # section headings so "## Summary" becomes Heading 1 in the document.
    levels = []
    title_seen = False
    for k, t in enumerate(tokens):
        if t.type != "heading_open":
            continue
        lvl = int(t.tag[1])
        text = tokens[k + 1].content.strip() if k + 1 < len(tokens) else ""
        if lvl == 1 and skip_title and not title_seen and _norm(text) == _norm(skip_title):
            title_seen = True
            continue
        levels.append(lvl)
    offset = (min(levels) - 1) if levels else 0
    lists: list[dict[str, Any]] = []
    quote = 0
    skipped_title = False
    i = 0
    n = len(tokens)
    while i < n:
        tok = tokens[i]
        kind = tok.type

        if kind == "heading_open":
            level = int(tok.tag[1])
            inline = tokens[i + 1]
            text = inline.content.strip()
            if (level == 1 and skip_title and not skipped_title
                    and _norm(text) == _norm(skip_title)):
                skipped_title = True
            else:
                name = {1: "Heading 1", 2: "Heading 2"}.get(max(1, level - offset), "Heading 3")
                _inline(doc.add_paragraph(style=_style(doc, name)), inline)
            i += 3
            continue

        if kind == "paragraph_open":
            inline = tokens[i + 1]
            content = inline.content.strip()
            if content in PAGEBREAK_MARKERS:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                i += 3
                continue
            if lists:
                ctx = lists[-1]
                level = len(lists)
                if ctx["first"]:
                    ctx["first"] = False
                    if ctx["ordered"]:
                        p = doc.add_paragraph(style=_style(doc, "List Paragraph"))
                        p.paragraph_format.left_indent = Cm(0.75 * level)
                        p.paragraph_format.first_line_indent = Cm(-0.6)
                        p.add_run(f"{ctx['n']}.  ")
                    else:
                        p = doc.add_paragraph(style=_style(doc, {1: "List Bullet", 2: "List Bullet 2"}.get(level, "List Bullet 3")))
                else:
                    p = doc.add_paragraph(style=_style(doc, "List Paragraph"))
                    p.paragraph_format.left_indent = Cm(0.75 * level)
                _inline(p, inline)
            elif quote:
                _inline(doc.add_paragraph(style=_style(doc, "Quote")), inline)
            else:
                _inline(doc.add_paragraph(style=_style(doc, "Normal")), inline)
            i += 3
            continue

        if kind in ("bullet_list_open", "ordered_list_open"):
            start = 1
            if kind == "ordered_list_open":
                try:
                    start = int(tok.attrs.get("start") or 1)
                except (TypeError, ValueError):
                    start = 1
            lists.append({"ordered": kind == "ordered_list_open", "n": start - 1, "first": False})
        elif kind in ("bullet_list_close", "ordered_list_close"):
            if lists:
                lists.pop()
        elif kind == "list_item_open":
            if lists:
                lists[-1]["n"] += 1
                lists[-1]["first"] = True
        elif kind == "blockquote_open":
            quote += 1
        elif kind == "blockquote_close":
            quote = max(0, quote - 1)
        elif kind in ("fence", "code_block"):
            p = doc.add_paragraph(style=_style(doc, "Normal"))
            p.paragraph_format.left_indent = Cm(0.5)
            lines = tok.content.rstrip("\n").split("\n")
            for j, line in enumerate(lines):
                run = p.add_run(line)
                run.font.name = MONO
                run.font.size = Pt(9)
                if j < len(lines) - 1:
                    run.add_break()
        elif kind == "hr":
            _border(doc.add_paragraph())
        elif kind == "table_open":
            rows: list[tuple[list[Any], bool]] = []
            in_head = False
            current: list[Any] | None = None
            j = i + 1
            while j < n and tokens[j].type != "table_close":
                t = tokens[j]
                if t.type == "thead_open":
                    in_head = True
                elif t.type == "thead_close":
                    in_head = False
                elif t.type == "tr_open":
                    current = []
                elif t.type == "tr_close" and current is not None:
                    rows.append((current, in_head))
                    current = None
                elif t.type in ("th_open", "td_open") and current is not None:
                    current.append(tokens[j + 1])
                j += 1
            _add_table(doc, rows)
            i = j + 1
            continue
        i += 1


def _add_table(doc, rows) -> None:
    from docx.shared import Pt

    if not rows:
        return
    cols = max(len(r) for r, _ in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    _table_borders(table)
    for r_idx, (cells, is_head) in enumerate(rows):
        row = table.rows[r_idx]
        if is_head:
            _repeat_header(row)
        for c_idx in range(cols):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if c_idx < len(cells):
                _inline(p, cells[c_idx], bold=is_head, size=Pt(9.5),
                        color=(0xFF, 0xFF, 0xFF) if is_head else None)
            if is_head:
                _shade_cell(cell, "%02X%02X%02X" % PINE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------------------
# DOCX: document assembly
# ---------------------------------------------------------------------------

def _cover(doc, meta: dict[str, Any]) -> None:
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    mark = doc.add_paragraph()
    run = mark.add_run("WIJERCO")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = _rgb(GOLD)
    mark.paragraph_format.space_after = Pt(0)

    title = doc.add_paragraph(meta.get("title") or "Untitled", style=_style(doc, "Title"))
    title.paragraph_format.space_before = Pt(150)
    if meta.get("type_label"):
        doc.add_paragraph(meta["type_label"], style=_style(doc, "Subtitle"))
    _border(doc.add_paragraph(), color="%02X%02X%02X" % GOLD, size=12)

    final = meta.get("status") == "final"
    lines = [
        ("Prepared for", meta.get("client")),
        ("Audience", meta.get("audience")),
        ("Date", meta.get("date") or date.today().strftime("%d %B %Y").lstrip("0")),
        ("Version", f"v{meta.get('version', 1)} · {'Approved' if final else 'Draft for review'}"),
    ]
    for label, value in lines:
        if not value:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        a = p.add_run(f"{label}   ")
        a.font.size = Pt(9.5)
        a.font.color.rgb = _rgb(MUTED)
        b = p.add_run(str(value))
        b.font.size = Pt(10.5)
    if not final:
        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(18)
        r = note.add_run("Draft for review. Not for circulation.")
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = _rgb(MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _header_footer(doc, meta: dict[str, Any]) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt

    final = meta.get("status") == "final"
    marker = "CONFIDENTIAL" if final else "DRAFT FOR REVIEW"
    for section in doc.sections:
        section.different_first_page_header_footer = True
        width = section.page_width - section.left_margin - section.right_margin
        head = section.header.paragraphs[0]
        head.text = ""
        head.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        a = head.add_run(f"WijerCo  ·  {meta.get('title') or ''}\t")
        a.font.size = Pt(8)
        a.font.color.rgb = _rgb(MUTED)
        b = head.add_run(marker)
        b.font.size = Pt(8)
        b.bold = True
        b.font.color.rgb = _rgb(PINE if final else GOLD)

        foot = section.footer.paragraphs[0]
        foot.text = ""
        foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        c = foot.add_run(f"{meta.get('type_label') or 'Deliverable'}  ·  v{meta.get('version', 1)}\tPage ")
        c.font.size = Pt(8)
        c.font.color.rgb = _rgb(MUTED)
        _field(foot, "PAGE", size=Pt(8), color=MUTED)
        d = foot.add_run(" of ")
        d.font.size = Pt(8)
        d.font.color.rgb = _rgb(MUTED)
        _field(foot, "NUMPAGES", size=Pt(8), color=MUTED)
        _ = Cm  # keep import for type checkers


def render_docx(markdown: str, out_path: Path | str, meta: dict[str, Any] | None = None) -> dict[str, Any]:
    """Render note markdown to a branded DOCX at out_path.

    meta: title, type_label, client, audience, date, version, status
    ("draft" or "final"), deliverable_id.
    """
    meta = dict(meta or {})
    out = Path(out_path)
    if out.exists():
        raise FileExistsError(f"refusing to overwrite {out}")
    document_md, _notes = split_working_notes(markdown)
    doc = _new_document()
    _cover(doc, meta)
    _header_footer(doc, meta)
    _add_markdown(doc, document_md, skip_title=meta.get("title"))

    props = doc.core_properties
    props.title = meta.get("title") or ""
    props.author = "WijerCo"
    props.subject = meta.get("type_label") or ""
    props.comments = f"Rendered from Deliverables note {meta.get('deliverable_id', '')} v{meta.get('version', 1)}"
    props.keywords = "WijerCo; deliverable"

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    words = word_count(document_md)
    return {
        "path": str(out),
        "bytes": out.stat().st_size,
        "words": words,
        "pages_est": 1 + max(1, round(words / 420)),
        "format": "docx",
    }


# ---------------------------------------------------------------------------
# PPTX: slide decks
# ---------------------------------------------------------------------------

_NOTES_RE = re.compile(r"^(?:>\s*)?(?:\*\*|__)?(?:speaker\s+)?notes?(?:\*\*|__)?\s*:\s*(?:\*\*|__)?\s*(.*)$", re.I)


def parse_slides(markdown: str) -> dict[str, Any]:
    """Read the deck note format.

        # Deck title
        Optional subtitle line

        ## Slide title
        - bullet
        Notes: what to say out loud
    """
    document_md, _ = split_working_notes(markdown)
    title, subtitle = "", ""
    slides: list[dict[str, Any]] = []
    cur: dict[str, Any] | None = None
    in_notes = False
    for raw in document_md.splitlines():
        s = raw.strip()
        if not s:
            continue
        if s.startswith("# ") and not title and cur is None:
            title = strip_inline_md(s[2:])
            continue
        if s.startswith("## "):
            heading = re.sub(r"^slide\s*\d+\s*[:.\-]\s*", "", strip_inline_md(s[3:]), flags=re.I)
            cur = {"title": heading, "bullets": [], "body": [], "notes": []}
            slides.append(cur)
            in_notes = False
            continue
        m = _NOTES_RE.match(s)
        if cur is None:
            if not subtitle and not m and not s.startswith(("-", "*", ">", "|")):
                subtitle = strip_inline_md(s)
            continue
        if m:
            in_notes = True
            if m.group(1).strip():
                cur["notes"].append(strip_inline_md(m.group(1)))
            continue
        if in_notes:
            cur["notes"].append(strip_inline_md(s.lstrip("> ")))
            continue
        bullet = re.match(r"^(?:[-*+]|\d+[.)])\s+(.*)$", s)
        if bullet:
            cur["bullets"].append(strip_inline_md(bullet.group(1)))
        elif not s.startswith("|"):
            cur["body"].append(strip_inline_md(s))
    return {"title": title, "subtitle": subtitle, "slides": slides}


def render_pptx(markdown: str, out_path: Path | str, meta: dict[str, Any] | None = None) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    meta = dict(meta or {})
    out = Path(out_path)
    if out.exists():
        raise FileExistsError(f"refusing to overwrite {out}")
    deck = parse_slides(markdown)
    title = deck["title"] or meta.get("title") or "Untitled deck"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    pine, gold, ink, muted = RGBColor(*PINE), RGBColor(*GOLD), RGBColor(*INK), RGBColor(*MUTED)
    white = RGBColor(0xFF, 0xFF, 0xFF)

    def rect(slide, x, y, w, h, color):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def text(slide, x, y, w, h, value, size, color, bold=False, align=None, anchor=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        if anchor is not None:
            tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = value
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return box

    total = len(deck["slides"]) + 1
    status = "Approved" if meta.get("status") == "final" else "Draft for review"
    footer_meta = "  ·  ".join(x for x in (meta.get("client"), meta.get("date"), status) if x)

    # Title slide
    s0 = prs.slides.add_slide(blank)
    rect(s0, 0, 0, W, H, pine)
    text(s0, Inches(0.8), Inches(0.7), Inches(6), Inches(0.5), "WIJERCO", 14, gold, bold=True)
    text(s0, Inches(0.8), Inches(2.3), Inches(11.6), Inches(2.2), title, 40, white, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    rect(s0, Inches(0.8), Inches(4.7), Inches(1.4), Inches(0.07), gold)
    if deck["subtitle"]:
        text(s0, Inches(0.8), Inches(4.95), Inches(11.6), Inches(0.9), deck["subtitle"], 20,
             RGBColor(0xE9, 0xD9, 0xB5))
    if footer_meta:
        text(s0, Inches(0.8), Inches(6.6), Inches(11.6), Inches(0.4), footer_meta, 12,
             RGBColor(0xB9, 0xC9, 0xC2))
    s0.notes_slide.notes_text_frame.text = (
        f"Open with the purpose of this deck: {title}." + (f" {deck['subtitle']}" if deck["subtitle"] else "")
    )

    for idx, sl in enumerate(deck["slides"], start=2):
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, Inches(0.18), H, pine)
        text(s, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.95), sl["title"],
             28 if len(sl["title"]) <= 48 else 24, pine, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        rect(s, Inches(0.7), Inches(1.5), Inches(1.2), Inches(0.06), gold)
        items = [("•", b) for b in sl["bullets"]] + [("", b) for b in sl["body"]]
        size = 20 if len(items) <= 4 else 18 if len(items) <= 6 else 15
        box = s.shapes.add_textbox(Inches(0.7), Inches(1.85), Inches(11.9), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        for k, (mark, value) in enumerate(items):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            r = p.add_run()
            r.text = f"{mark}  {value}" if mark else value
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = ink
        text(s, Inches(0.7), Inches(6.95), Inches(6), Inches(0.35), "WijerCo", 11, muted)
        text(s, Inches(6.9), Inches(6.95), Inches(5.7), Inches(0.35), f"{idx} / {total}", 11, muted,
             align=PP_ALIGN.RIGHT)
        notes = " ".join(sl["notes"]).strip()
        if not notes:
            points = "; ".join(sl["bullets"][:5])
            notes = f"Explain {sl['title']}." + (f" Cover: {points}." if points else "")
        s.notes_slide.notes_text_frame.text = notes

    props = prs.core_properties
    props.title = title
    props.author = "WijerCo"
    props.subject = meta.get("type_label") or "Slide deck"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return {
        "path": str(out),
        "bytes": out.stat().st_size,
        "slides": total,
        "format": "pptx",
    }


def render(markdown: str, out_path: Path | str, meta: dict[str, Any] | None = None) -> dict[str, Any]:
    """Dispatch on the output file's extension."""
    suffix = Path(out_path).suffix.lower()
    if suffix == ".pptx":
        return render_pptx(markdown, out_path, meta)
    if suffix == ".docx":
        return render_docx(markdown, out_path, meta)
    raise ValueError(f"unsupported output type: {suffix}")
